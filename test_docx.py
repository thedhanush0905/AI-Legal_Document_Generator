import json
import os
import sys
import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure UTF-8 stdout on Windows
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

from src.affidavit_models import GeneratedAffidavit
from src.docx_generator import create_affidavit_docx


def main():
    print("========================================")
    print("STEP 5: DOCX GENERATION TEST")
    print("========================================\n")

    input_json_path = os.path.join("outputs", "generated_affidavit.json")
    output_docx_path = os.path.join("outputs", "generated_affidavit.docx")

    # 1. Load outputs/generated_affidavit.json
    print(f"📂 1. Loading GeneratedAffidavit from {input_json_path}...")
    with open(input_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # 2. Validate it using GeneratedAffidavit Pydantic model
    print("🔍 2. Validating GeneratedAffidavit model with Pydantic...")
    affidavit = GeneratedAffidavit.model_validate(data)
    print("   ✅ GeneratedAffidavit validated successfully.")

    # 3. Call create_affidavit_docx()
    print(f"\n📄 3. Creating DOCX file at {output_docx_path}...")
    abs_path = create_affidavit_docx(affidavit, output_docx_path)
    print(f"   ✅ DOCX created at: {abs_path}")

    # 4. Verify DOCX file exists
    assert os.path.exists(output_docx_path), f"File not found: {output_docx_path}"
    file_size = os.path.getsize(output_docx_path)
    print(f"   ✅ File exists on disk (Size: {file_size} bytes).")

    # 5. Open resulting DOCX using python-docx and inspect text & formatting
    print("\n🔍 4. Inspecting DOCX contents and formatting...")
    doc = Document(output_docx_path)

    # Extract all text from paragraphs and tables
    extracted_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_texts = []
    for t in doc.tables:
        for r in t.rows:
            row_cells = [c.text.strip() for c in r.cells if c.text.strip()]
            table_texts.append(" | ".join(row_cells))

    full_document_text = "\n".join(extracted_paragraphs + table_texts)

    # 6. Assert required case entities are present
    required_strings = [
        "IN THE HIGH COURT OF JUDICATURE AT BOMBAY",
        "ORDINARY ORIGINAL CIVIL JURISDICTION",
        "WRIT PETITION NO. 1847 OF 2026",
        "Sunrise Housing Private Limited",
        "State of Maharashtra",
        "Mumbai Metropolitan Region Development Authority",
        "AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. 2",
        "Arvind Rajan",
        "Deputy Metropolitan Commissioner",
        "EXHIBIT-‘A’",
        "PRAYER",
        "dismiss the present Writ Petition with costs",
        "Solemnly affirmed at Mumbai",
        "5th day of September 2026",
        "VERIFICATION",
        "paragraphs 1 to 7",
        "Rajan & Associates",
    ]

    print("\n📋 5. Checking presence of mandatory textual content:")
    for text_item in required_strings:
        assert text_item in full_document_text, f"Missing required text: '{text_item}'"
        print(f"   ✓ '{text_item}'")

    # 7. Check body paragraph count and prayer
    body_paras = [p for p in doc.paragraphs if any(p.text.strip().startswith(f"{i}. ") for i in range(1, 8))]
    assert len(body_paras) == 7, f"Expected exactly 7 body paragraphs in DOCX, found {len(body_paras)}"
    print(f"   ✓ Exactly {len(body_paras)} body paragraphs rendered.")

    # Assert no placeholder text exists in any paragraph
    placeholder_indicators = ["paragraph 4 text", "paragraph 5 text", "paragraph 6 text", "placeholder", "to be drafted"]
    for bp in body_paras:
        for indicator in placeholder_indicators:
            assert indicator not in bp.text.lower(), f"Placeholder text found in paragraph: '{bp.text}'"
    print("   ✓ No placeholder or description prefixes present in body paragraphs.")

    assert "(a)" in full_document_text, "Missing '(a)' in prayer"
    print("   ✓ Prayer letter '(a)' is present.")

    # 8. Hallucination check: ensure no sample facts leaked
    prohibited_sample_entities = ["Arjun Mehta", "Rohan Deshpande", "3147"]
    for prohibited in prohibited_sample_entities:
        assert prohibited not in full_document_text, f"Sample fact '{prohibited}' leaked into DOCX!"
    print("   ✓ Hallucination check passed: No sample facts present in DOCX.")

    # 9. Formatting assertions
    print("\n🎨 6. Validating document formatting rules:")

    # Check headings (Forum heading, Jurisdiction, Case number, Affidavit title, PRAYER, VERIFICATION)
    headings_to_check = [
        affidavit.forum_heading,
        affidavit.jurisdiction,
        affidavit.case_number_line,
        affidavit.affidavit_title,
        "PRAYER",
        "VERIFICATION",
    ]

    for heading_text in headings_to_check:
        matching_p = next(
            (p for p in doc.paragraphs if p.text.strip().upper() == heading_text.upper()),
            None,
        )
        assert matching_p is not None, f"Heading not found: {heading_text}"
        assert matching_p.alignment == WD_ALIGN_PARAGRAPH.CENTER, f"Heading '{heading_text}' is not centered"
        # Check that runs in heading are bold
        assert any(r.bold for r in matching_p.runs), f"Heading '{heading_text}' is not bold"
        print(f"   ✓ Heading '{heading_text}' is centered and bold.")

    # Check that body paragraphs are justified and paragraph numbers are bold
    for bp in body_paras:
        assert bp.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, f"Paragraph '{bp.text[:30]}...' is not justified"
        first_run = bp.runs[0]
        assert first_run.bold is True, f"Paragraph number for '{bp.text[:30]}...' is not bold"
    print("   ✓ All 7 body paragraphs are justified and have bold paragraph numbers.")

    # Check DEPONENT tag in tables
    deponent_runs = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "DEPONENT" in p.text:
                        assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT, "DEPONENT tag is not right-aligned"
                        assert p.text.strip() == "DEPONENT", "DEPONENT tag is not uppercase"
                        deponent_runs.append(p)
    assert len(deponent_runs) >= 2, f"Expected at least 2 DEPONENT signatures (jurat + verification), found {len(deponent_runs)}"
    print(f"   ✓ DEPONENT tags ({len(deponent_runs)}) are uppercase and right-aligned.")

    # Check Before Me is left aligned
    before_me_found = False
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "Before Me" in p.text:
                        assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT, "'Before Me' is not left-aligned"
                        before_me_found = True
    assert before_me_found, "'Before Me' not found in jurat table"
    print("   ✓ 'Before Me' is left-aligned in Jurat.")

    print("\n========================================")
    print("STEP 5 TEST PASSED SUCCESSFULLY!")
    print("========================================")


if __name__ == "__main__":
    main()
