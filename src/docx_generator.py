"""
DOCX Generator for Affidavit in Reply.
Renders a validated GeneratedAffidavit into a properly formatted .docx file
using python-docx adhering to the High Court of Bombay legal formatting conventions.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from src.affidavit_models import GeneratedAffidavit


def create_affidavit_docx(affidavit: GeneratedAffidavit, output_path: str) -> str:
    """
    Renders the structured GeneratedAffidavit into a DOCX document.

    Args:
        affidavit (GeneratedAffidavit): The structured affidavit data model.
        output_path (str): Target filesystem path for the .docx file.

    Returns:
        str: Absolute path to the created .docx file.
    """
    doc = Document()

    # Configure 1-inch margins
    sections = doc.sections
    for sec in sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Set default style font: Times New Roman, 12pt, black text
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # Helper for adding centered bold heading
    def add_centered_heading(text: str, space_before: int = 0, space_after: int = 6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text.upper())
        run.bold = True
        return p

    # 1. Forum Heading
    add_centered_heading(affidavit.forum_heading, space_before=0, space_after=4)

    # 2. Jurisdiction
    add_centered_heading(affidavit.jurisdiction, space_before=2, space_after=4)

    # 3. Case Number Line
    add_centered_heading(affidavit.case_number_line, space_before=2, space_after=12)

    # 4. Cause Title: Rendered cleanly using a borderless 2-column table
    # Left column: Party name/details, Right column: Right-aligned status tag (...Petitioner)
    ct = affidavit.cause_title
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Set column widths (e.g. 4.5 inches left, 2.0 inches right)
    col_widths = [Inches(4.5), Inches(2.0)]

    def add_party_row(party_name: str, status_tag: str):
        row = table.add_row()
        # Left cell
        cell_left = row.cells[0]
        cell_left.width = col_widths[0]
        p_l = cell_left.paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_l.paragraph_format.space_after = Pt(2)
        p_l.paragraph_format.space_before = Pt(2)
        p_l.add_run(party_name)

        # Right cell
        cell_right = row.cells[1]
        cell_right.width = col_widths[1]
        p_r = cell_right.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_r.paragraph_format.space_after = Pt(2)
        p_r.paragraph_format.space_before = Pt(2)
        p_r.add_run(status_tag)

    def add_versus_row(versus_text: str = "VERSUS"):
        row = table.add_row()
        # Merge cells across table
        cell = row.cells[0]
        cell.merge(row.cells[1])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(versus_text)
        run.bold = True

    # Petitioner
    add_party_row(ct.petitioner.name, ct.petitioner.status_tag)
    # VERSUS
    add_versus_row(ct.versus_text)
    # Respondents
    for resp in ct.respondents:
        add_party_row(resp.name, resp.status_tag)

    # Spacing after cause title
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(6)
    spacer.paragraph_format.space_after = Pt(6)

    # 5. Affidavit Title
    add_centered_heading(affidavit.affidavit_title, space_before=8, space_after=12)

    # 6. Deponent Clause
    p_dep = doc.add_paragraph()
    p_dep.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dep.paragraph_format.space_after = Pt(8)
    p_dep.paragraph_format.line_spacing = 1.15
    p_dep.add_run(affidavit.deponent_clause)

    # 7. Numbered Reply Paragraphs
    for bp in affidavit.body_paragraphs:
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.space_after = Pt(6)
        p_body.paragraph_format.line_spacing = 1.15

        # Bold paragraph number
        run_num = p_body.add_run(f"{bp.paragraph_number}. ")
        run_num.bold = True

        # Normal body text
        p_body.add_run(bp.text)

    # 8. Prayer Section
    add_centered_heading(affidavit.prayer.heading, space_before=12, space_after=6)

    p_prayer_intro = doc.add_paragraph()
    p_prayer_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_prayer_intro.paragraph_format.space_after = Pt(4)
    p_prayer_intro.add_run(affidavit.prayer.intro_text)

    for pi in affidavit.prayer.items:
        p_item = doc.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(4)

        # Bold letter tag e.g. (a)
        run_letter = p_item.add_run(f"{pi.letter} ")
        run_letter.bold = True

        p_item.add_run(pi.text)

    # 9. Jurat Section
    p_jurat_place = doc.add_paragraph()
    p_jurat_place.paragraph_format.space_before = Pt(12)
    p_jurat_place.paragraph_format.space_after = Pt(2)
    p_jurat_place.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_jurat_place.add_run(affidavit.jurat.place_line)

    p_jurat_date = doc.add_paragraph()
    p_jurat_date.paragraph_format.space_after = Pt(4)
    p_jurat_date.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_jurat_date.add_run(affidavit.jurat.date_line)

    # Jurat signature line: Before Me on left, DEPONENT on right
    jurat_table = doc.add_table(rows=1, cols=2)
    jurat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    jurat_table.autofit = False
    jurat_table.rows[0].cells[0].width = Inches(3.25)
    jurat_table.rows[0].cells[1].width = Inches(3.25)

    p_bm = jurat_table.rows[0].cells[0].paragraphs[0]
    p_bm.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_bm.add_run(affidavit.jurat.before_me)

    p_dep_sig = jurat_table.rows[0].cells[1].paragraphs[0]
    p_dep_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_dep = p_dep_sig.add_run(affidavit.jurat.deponent_tag.upper())
    run_dep.bold = True

    # 10. Verification Section
    add_centered_heading(affidavit.verification.heading, space_before=14, space_after=6)

    p_ver_body = doc.add_paragraph()
    p_ver_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ver_body.paragraph_format.space_after = Pt(6)
    p_ver_body.paragraph_format.line_spacing = 1.15
    p_ver_body.add_run(affidavit.verification.full_text)

    # Verification place & date + DEPONENT right aligned
    ver_table = doc.add_table(rows=1, cols=2)
    ver_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ver_table.autofit = False
    ver_table.rows[0].cells[0].width = Inches(4.5)
    ver_table.rows[0].cells[1].width = Inches(2.0)

    p_vat = ver_table.rows[0].cells[0].paragraphs[0]
    p_vat.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_vat.add_run(affidavit.verification.verified_at_line)

    p_vdep = ver_table.rows[0].cells[1].paragraphs[0]
    p_vdep.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_vdep = p_vdep.add_run(affidavit.verification.deponent_tag.upper())
    run_vdep.bold = True

    # 11. Advocate Block (foot of document)
    p_adv_firm = doc.add_paragraph()
    p_adv_firm.paragraph_format.space_before = Pt(14)
    p_adv_firm.paragraph_format.space_after = Pt(2)
    p_adv_firm.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_firm = p_adv_firm.add_run(affidavit.advocate_block.firm_name)
    run_firm.bold = True

    p_adv_acting = doc.add_paragraph()
    p_adv_acting.paragraph_format.space_after = Pt(6)
    p_adv_acting.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_adv_acting.add_run(affidavit.advocate_block.acting_for)

    # Ensure parent directories exist and save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return os.path.abspath(output_path)
